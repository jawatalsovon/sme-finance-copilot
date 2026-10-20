"""bd_grocer exports: bKash-style wallet CSV, bank CSV, Google-Sheets-style notebook xlsx, canonical tables."""
from __future__ import annotations

import datetime as dt
from collections import defaultdict
from pathlib import Path

import numpy as np
from openpyxl import Workbook

import canon
import bd
from core import Trace, ym, money, month_end, rng_for, save_xlsx, write_csv, D0, D1, add_months, days

CID = bd.CID
BOOK_SYSTEM = "gsheets"
BOOK_FILE = "notebook_<year>.xlsx"


def system_for(file: str) -> str:
    return "gsheets"


def fx_at(A):
    return lambda ccy, d: 1.0
WALLET_NO = "01700-000451"
BANK_NO = "****4821"
MODE_LABEL = {"cash": "Cash", "bank": "Bank", "wallet": "bKash", "ap": "Udhar", "ar": "Baki", "adv": "Dadon", "inv": "Stock"}
MODE_PRIORITY = ["cash", "bank", "wallet", "ap", "ar", "adv", "inv"]
HEAD_LABEL = {"sales": "Sales", "cogs": "COGS (stock count)", "shrink": "Stock loss", "sal": "Salary", "rent": "Rent", "util": "Utilities",
              "deliv": "Delivery", "fees": "Charges", "repair": "Repair", "misc": "Misc", "lic": "Licence", "intr": "Loan interest",
              "bad": "Bad debt", "draw": "Drawings", "cap": "Capital", "loan": "Loan principal", "inv": "Purchase", "adv": "Dadon",
              "ar": "Baki collection", "fa": "Equipment", "othinc": "Other income", "ap": "Supplier payment", "bank": "Transfer from bank",
              "cash": "Transfer from cash", "wallet": "Transfer from bKash"}
BANGLA = str.maketrans("0123456789", "০১২৩৪৫৬৭৮৯")


def mode_and_head(e):
    """For a 2-line entry: (mode_acct, head_acct, direction 'In'/'Out', amount)."""
    keys = [l.acct for l in e.lines]
    mode = next((k for k in MODE_PRIORITY if k in keys), keys[0])
    ml = next(l for l in e.lines if l.acct == mode)
    head = next(l.acct for l in e.lines if l is not ml)
    return mode, head, ("In" if ml.dr else "Out"), (ml.dr or ml.cr)


def bank_narration(e, sim, by_bill_entry):
    k = e.kind
    sup = None
    if k in ("purchase_bank", "bill_payment", "advance_paid"):
        s = sim.entry_info.get(e.id, {}).get("sup")
        if s is None and e.id in by_bill_entry:
            s = by_bill_entry[e.id]
        sup = bd.WHOLESALERS[s] if s is not None else "WHOLESALER"
    return {
        "cash_deposit": "CASH DEPOSIT BY SELF", "wallet_to_bank": "NPSB BKASH MERCHANT SETTLEMENT", "sale_card": f"POS SETTLEMENT {e.date:%d%m%y}",
        "purchase_bank": f"CHQ PAID {sup}", "bill_payment": f"CHQ PAID {sup}", "advance_paid": f"CHQ ADVANCE {sup}", "rent": "CHQ PAID SHOP RENT",
        "utilities": "DPDC ELECTRICITY BILL PAYMENT", "salary_bank": "SALARY TRF - MANAGER", "drawings": "CHQ SELF", "bank_fee": "ACCOUNT MAINT AND CHQ BOOK CHARGES",
        "loan_payment": "LOAN INSTALMENT RECOVERY", "owner_injection": "CASH DEPOSIT - PROPRIETOR", "licence": "CHQ PAID LICENCE AND SIGNBOARD TAX",
        "equipment": "CHQ PAID EQUIPMENT SUPPLIER", "loan_drawdown": "LOAN DISBURSEMENT", "opening": "BALANCE B/F", "misc_cash": "ATM CASH WITHDRAWAL",
        "salary_cash": "ATM CASH WITHDRAWAL", "repair": "CHQ PAID REPAIR", "delivery": "ATM CASH WITHDRAWAL",
    }.get(k, "BANK TRANSACTION")


def build_statements(sim, seed):
    """Reality: wallet and bank statement rows (with running balances) and monthly cash counts, from the TRUTH entries."""
    w = sim.world
    rng = rng_for(seed, CID, "stmt")
    by_ref = {e.ref: e for e in w.entries if e.kind == "wallet_fee"}
    by_grp = defaultdict(list)
    for e in w.entries:
        if e.tags.get("grp"):
            by_grp[e.tags["grp"]].append(e)
    bills_by_id = {b["id"]: b for b in sim.bills}
    bill_by_entry = {p["entry"]: bills_by_id[p["bill"]]["sup"] for p in sim.bill_payments}
    for b in sim.bills:
        bill_by_entry[b["entry"]] = b["sup"]
    rows = {"wallet": [], "bank": []}
    # ---- wallet
    txn_by_baki = {}
    for d, txs in sim.wallet_txns.items():
        for t in txs:
            if "baki" in t:
                txn_by_baki[t["baki"]] = t
    sale_wallet = {e.date: e for e in w.entries if e.kind == "sale_wallet"}
    day_events: dict[dt.date, list[dict]] = defaultdict(list)
    for e in w.entries:
        if e.kind == "opening":
            for l in e.lines:
                if l.acct in ("wallet", "bank"):
                    rows[l.acct].append(dict(date=e.date, tm="00:00", desc="Opening balance brought forward", cp="", ref="OPENING", amount=l.nat, fee=0, entries=[e.id], kind="opening", ext=f"OPEN{l.acct}"))
            continue
        if e.kind == "sale_wallet":
            txs = [t for t in sim.wallet_txns[e.date] if "baki" not in t]
            fe = by_ref[f"F{e.date:%Y%m%d}"]
            for j, t in enumerate(txs):
                day_events[e.date].append(dict(acct="wallet", desc="Payment received", cp=t["cp"], ref=f"INV{e.date:%d%m}{j:03d}", amount=t["amount"] - t["fee"], fee=t["fee"],
                                               gross=t["amount"], entries=[e.id, fe.id], kind="sale_wallet", credit=True))
        elif e.kind == "baki_wallet":
            t = txn_by_baki[e.id]
            fe = by_ref[f"F{e.date:%Y%m%d}-{e.tags['hh']}"]
            day_events[e.date].append(dict(acct="wallet", desc="Payment received", cp=t["cp"], ref="", amount=t["amount"] - t["fee"], fee=t["fee"], gross=t["amount"], entries=[e.id, fe.id], kind="baki_wallet", credit=True))
        elif e.kind == "purchase_wallet":
            fe = next(x for x in by_grp[e.tags["grp"]] if x.kind == "wallet_fee")
            wl = next(l for l in e.lines if l.acct == "wallet")
            sup = bd.WHOLESALERS[sim.entry_info[e.id]["sup"]]
            day_events[e.date].append(dict(acct="wallet", desc="Send Money", cp=sup, ref=e.ref, amount=-(wl.cr + fe.lines[0].dr), fee=fe.lines[0].dr, gross=wl.cr, entries=[e.id, fe.id], kind="purchase_wallet", credit=False))
        elif e.kind in ("personal_wallet", "wallet_to_bank") or (e.kind == "utilities" and any(l.acct == "wallet" for l in e.lines)):
            wl = next(l for l in e.lines if l.acct == "wallet")
            desc = {"personal_wallet": "Send Money", "wallet_to_bank": "Transfer to Bank", "utilities": "Bill Payment"}[e.kind]
            cp = {"personal_wallet": "01XXXXXXXXX", "wallet_to_bank": "Mirpur Commercial Bank " + BANK_NO, "utilities": "Utility biller"}[e.kind]
            day_events[e.date].append(dict(acct="wallet", desc=desc, cp=cp, ref=e.ref, amount=-wl.cr, fee=0, gross=wl.cr, entries=[e.id], kind=e.kind, credit=False))
    for d in sorted(day_events):
        ev = day_events[d]
        cr = [x for x in ev if x["credit"]]
        dr = [x for x in ev if not x["credit"]]
        cts = sorted(rng.uniform(8 * 60, 21.5 * 60, size=len(cr)))
        dts = sorted(rng.uniform(21.6 * 60, 23.9 * 60, size=len(dr)))
        for x, m in list(zip(cr, cts)) + list(zip(dr, dts)):
            x["date"], x["tm"] = d, f"{int(m // 60):02d}:{int(m % 60):02d}"
            x["ext"] = "".join("ABCDEFGHJKLMNPQRSTUVWXYZ0123456789"[int(i)] for i in rng.integers(0, 34, size=10))
            rows["wallet"].append(x)
    rows["wallet"].sort(key=lambda x: (x["date"], x["tm"]) if x["kind"] != "opening" else (x["date"], "00:00"))
    # ---- bank (one row per date + group)
    grp: dict[tuple, dict] = {}
    for e in w.entries:
        for l in e.lines:
            if l.acct != "bank":
                continue
            if e.kind == "opening":
                continue
            key = (e.date, e.tags.get("grp") or e.id)
            g = grp.setdefault(key, dict(date=e.date, desc=bank_narration(e, sim, bill_by_entry), cp="", ref=e.ref, amount=0, fee=0, entries=[], kind=e.kind, ext=""))
            g["amount"] += l.nat
            g["entries"].append(e.id)
            if e.kind == "sale_card":
                g["kind"] = "sale_card"
    for k in sorted(grp):
        rows["bank"].append(grp[k])
    for i, r in enumerate(rows["bank"]):
        if r["kind"] != "opening":
            r["ext"] = f"BNK{r['date']:%y%m%d}{i % 10000:04d}"
    # running balances
    for acct in rows:
        bal = 0
        for r in rows[acct]:
            bal += r["amount"]
            r["balance"] = bal
        assert bal == w.bal(acct), (acct, bal, w.bal(acct))
    # ---- monthly cash counts
    counts = []
    bal = 0
    cash_by_day: dict[dt.date, int] = defaultdict(int)
    for e in w.entries:
        for l in e.lines:
            if l.acct == "cash":
                cash_by_day[e.date] += l.nat
    for d in days():
        bal += cash_by_day.get(d, 0)
        if d == month_end(d):
            counts.append(dict(date=d, amount=bal))
    return rows, counts


def fmt_bank_date(d):
    return d.strftime("%d-%b-%Y")


def write_vendor_statements(sim, rows, counts, outdir: Path, trace: Trace):
    src = outdir / "sources"
    by_year: dict[int, list] = defaultdict(list)
    for r in rows["wallet"]:
        by_year[r["date"].year].append(r)
    for y, rs in sorted(by_year.items()):
        fn = f"bkash_style_wallet_statement_{y}.csv"
        out = []
        for i, r in enumerate(rs, 2):
            amt = r["amount"]
            cr = money(r["gross"]) if (amt > 0 and r["kind"] != "opening") else (money(amt) if r["kind"] == "opening" else "")
            db = money(r["gross"]) if amt < 0 and r["kind"] != "opening" else ""
            cp = r["cp"]
            if cp and cp[:2] == "01" and "X" not in cp:
                cp = cp[:5] + "XXX" + cp[-3:]
            out.append([r["ext"], r["date"].strftime("%d/%m/%Y"), r["tm"], r["desc"], cp, r["ref"], cr, db, money(r["fee"]) if r["fee"] else "0.00", money(r["balance"])])
            r.setdefault("src", []).append((fn, str(i)))
            for eid in r["entries"]:
                trace.add(fn, str(i), "truth", eid)
        write_csv(src / fn, ["TrxID", "Date", "Time", "Transaction Type", "Counterparty", "Reference", "Credit (BDT)", "Debit (BDT)", "Fee (BDT)", "Balance (BDT)"], out)
    by_year = defaultdict(list)
    for r in rows["bank"]:
        by_year[r["date"].year].append(r)
    for y, rs in sorted(by_year.items()):
        fn = f"mirpur_commercial_bank_statement_{y}.csv"
        out = []
        for i, r in enumerate(rs, 2):
            a = r["amount"]
            out.append([fmt_bank_date(r["date"]), fmt_bank_date(r["date"]), r["desc"], r["ref"] if r["desc"].startswith("CHQ") else "", money(-a) if a < 0 else "", money(a) if a > 0 else "", money(r["balance"])])
            r.setdefault("src", []).append((fn, str(i)))
            for eid in r["entries"]:
                trace.add(fn, str(i), "truth", eid)
        write_csv(src / fn, ["Txn Date", "Value Date", "Narration", "Chq/Ref No", "Withdrawal", "Deposit", "Balance"], out)


def statement_canonical(sim, rows, counts, T: canon.Tables, trace: Trace):
    w = sim.world
    n = 0
    for acct, chan, ref in (("wallet", "wallet", "01700-XXX451"), ("bank", "bank", BANK_NO)):
        for r in rows[acct]:
            n += 1
            f, rid = r["src"][0] if r.get("src") else ("", "")
            T.add("external_statement_lines", id=f"ext-{CID}-{n:06d}", company_id=CID, channel=chan, account_id=canon.acct_id(CID, w, acct), statement_ref=ref,
                  txn_date=r["date"], value_date=r["date"], amount=money(r["amount"]), currency="BDT", fx_rate="1.00000000", amount_base=money(r["amount"]),
                  description=r["desc"], counterparty_text=r["cp"] or None, balance=money(r["balance"]), external_txn_id=r["ext"], fee=money(r["fee"]),
                  source_system="wallet_csv" if acct == "wallet" else "bank_csv", source_file=f, source_row_id=rid)
            r["canon_id"] = f"ext-{CID}-{n:06d}"
    for ci, c in enumerate(counts, 2):
        n += 1
        T.add("external_statement_lines", id=f"ext-{CID}-{n:06d}", company_id=CID, channel="cash_count", account_id=canon.acct_id(CID, w, "cash"), statement_ref="till+safe",
              txn_date=c["date"], value_date=c["date"], amount=money(c["amount"]), currency="BDT", fx_rate="1.00000000", amount_base=money(c["amount"]),
              description="Month-end physical cash count (till and safe)", balance=money(c["amount"]), external_txn_id=f"COUNT{c['date']:%Y%m%d}", fee="0.00",
              source_system="cash_count", source_file="cash_counts.csv", source_row_id=str(ci))
        trace.add("cash_counts.csv", str(ci), "truth_balance", f"cash:{c['date']}")


# --------------------------------------------------------------------------- notebook
def _cell_date(d, rng, year):
    r = rng.random()
    if r < 0.72:
        return dt.datetime(d.year, d.month, d.day)
    if r < 0.84:
        return d.strftime("%d/%m/%Y")
    if r < 0.92:
        return f"{d.day}-{d.strftime('%b')}"
    if r < 0.97:
        return d.strftime("%m/%d/%Y") if d.day <= 12 else d.strftime("%d/%m/%Y")
    return d.strftime("%d.%m.%y")


def _cell_amount(cents, rng):
    taka = cents / 100
    whole = int(taka) if cents % 100 == 0 else taka
    r = rng.random()
    if r < 0.86:
        return whole
    if r < 0.94:
        return f"{taka:,.0f}" if cents % 100 == 0 else f"{taka:,.2f}"
    if r < 0.98:
        return f"Tk {taka:.0f}" if cents % 100 == 0 else f"Tk {taka:.2f}"
    if r < 0.995:
        return (f"{taka:.0f}" if cents % 100 == 0 else f"{taka:.2f}").translate(BANGLA)
    return f"{taka:.0f}/-"


TYPOS = {"Purchase": ["Purchse", "Purchace"], "Sales": ["Sale", "sales"], "Salary": ["Salry", "salary"], "Misc": ["Misc.", "misc"], "Utilities": ["Utility", "Electric"],
         "Rent": ["rent", "Shop rent"], "Delivery": ["Delivary"], "Charges": ["Charge", "bkash charge"]}


def write_notebook(sim, seed, outdir: Path, trace: Trace):
    rng = rng_for(seed, CID, "notebook")
    b = sim.book
    w = sim.world
    entries = b.sorted_entries()
    by_month: dict[str, list] = defaultdict(list)
    for e in entries:
        by_month[ym(e.date)].append(e)
    hhs = sim.households
    book_ids = set(b.entries)
    paid_ok = [p for p in sim.hh_payments if b.bid(p["entry"]) in book_ids]
    mapping = {"columns": ["Date", "Details", "Head", "Mode", "In", "Out", "Note"], "mode_labels": MODE_LABEL, "head_labels": HEAD_LABEL,
               "rule": "Row = one 2-line journal entry. Mode = the cash/bank/wallet (else Udhar/Baki/Dadon/Stock) side; In = Mode account debited, Out = credited; Head = the other account. Opening balances: Head 'Opening balance', one row per account."}
    for y in range(2020, 2027):
        wb = Workbook()
        wb.remove(wb.active)
        fn = f"notebook_{y}.xlsx"
        for m in range(1, 13):
            key = f"{y}-{m:02d}"
            if key not in by_month:
                continue
            ws = wb.create_sheet(dt.date(y, m, 1).strftime("%b%y"))
            ws["A1"] = f"Mirpur Family Store - Hisab {dt.date(y, m, 1):%B %Y}"
            ws.append([])
            ws.append(["Date", "Details", "Head", "Mode", "In", "Out", "Note"])
            rowno = 3
            late = []
            ordered = []
            for e in by_month[key]:
                (late if rng.random() < 0.03 else ordered).append(e)
            wk_in = wk_out = 0
            prev_week = None
            for e in ordered + late:
                if e.kind == "opening":
                    for l in e.lines:
                        rowno += 1
                        ws.append([_cell_date(e.date, rng, y), "Opening balance 2020", "Opening balance", MODE_LABEL.get(l.acct, HEAD_LABEL.get(l.acct, l.acct)),
                                   _cell_amount(l.dr, rng) if l.dr else None, _cell_amount(l.cr, rng) if l.cr else None, f"acct:{l.acct}"])
                        trace.add(fn, f"{ws.title}!R{rowno}", "book", e.id)
                    continue
                mode, head, direction, amt = mode_and_head(e)
                if rng.random() < 0.012 and e.kind in ("sale_cash", "misc_cash", "delivery"):
                    ws.append([])
                    rowno += 1
                    trace.add(fn, f"{ws.title}!R{rowno}", "layout", "blank_row")
                hlabel = HEAD_LABEL[head]
                if mode in ("inv",) and head == "cogs":
                    hlabel = "COGS (stock count)"
                if rng.random() < 0.035 and hlabel in TYPOS:
                    hlabel = TYPOS[hlabel][int(rng.integers(len(TYPOS[hlabel])))]
                if head in ("bank", "cash", "wallet"):
                    hlabel = f"Transfer {'to' if direction == 'Out' else 'from'} {MODE_LABEL[head]}"
                ml = MODE_LABEL[mode]
                r = rng.random()
                ml = ml if r < 0.9 else (ml.lower() if r < 0.95 else ml.upper() + " ")
                details = e.memo
                note = e.ref if (e.ref and rng.random() < 0.5) else None
                rowno += 1
                ws.append([_cell_date(e.date, rng, y), details, hlabel, ml, _cell_amount(amt, rng) if direction == "In" else None,
                           _cell_amount(amt, rng) if direction == "Out" else None, note])
                trace.add(fn, f"{ws.title}!R{rowno}", "book", e.id)
                wk = e.date.isocalendar()[1]
                if prev_week is not None and wk != prev_week and rng.random() < 0.5:
                    pass
                prev_week = wk
            ws.column_dimensions["B"].width = 44
        ws = wb.create_sheet("Baki")
        ws["A1"] = f"Baki (credit) book {y}"
        ws.append([])
        ws.append(["Date", "Household", "Particulars", "Given (credit)", "Received", "How paid"])
        rowno = 3
        rows_b = []
        for (d, hi), amt in sim.credit_sales.items():
            if d.year == y:
                rows_b.append((d, 0, hhs[hi].name, "Goods on credit", amt, None, None, sim.credit_entry[d]))
        for p in paid_ok:
            if p["date"].year == y:
                rows_b.append((p["date"], 1, hhs[p["hh"]].name, "Payment received", None, p["amount"], p["method"], p["entry"]))
        rows_b.sort(key=lambda x: (x[0], x[1], x[2]))
        for d, _, name, part, given, recv, how, eid in rows_b:
            rowno += 1
            ws.append([_cell_date(d, rng, y), name, part, _cell_amount(given, rng) if given else None, _cell_amount(recv, rng) if recv else None, how])
            trace.add(fn, f"Baki!R{rowno}", "truth", eid)
        save_xlsx(wb, outdir / "sources" / f"notebook_{y}.xlsx")
    from core import write_json
    write_json(outdir / "sources" / "notebook_mapping.json", mapping)


# --------------------------------------------------------------------------- canonical subledgers
def subledgers(sim, T_truth: canon.Tables, T_books: canon.Tables, trace: Trace, A):
    w = sim.world
    b = sim.book
    book_ids = set(b.entries)
    hhs = sim.households
    # counterparties
    for h in hhs:
        for T in (T_truth, T_books):
            T.add("counterparties", id=f"{CID}-{h.cid}", company_id=CID, kind="customer", name=h.name, country="BD", payment_terms_days=30, external_ref=h.cid,
                  source_system="synthetic" if T is T_truth else "gsheets", source_file="truth" if T is T_truth else "notebook (Baki tab)", source_row_id=h.cid)
    for i, s_ in enumerate(bd.WHOLESALERS):
        for T in (T_truth, T_books):
            T.add("counterparties", id=f"{CID}-sup-{i + 1:02d}", company_id=CID, kind="supplier", name=s_, country="BD", payment_terms_days=14, external_ref=f"sup-{i + 1:02d}",
                  source_system="synthetic" if T is T_truth else "gsheets", source_file="truth" if T is T_truth else "notebook", source_row_id=f"sup-{i + 1:02d}")
    for i, (role, _) in enumerate(bd.EMPLOYEES):
        for T in (T_truth, T_books):
            T.add("counterparties", id=f"{CID}-emp-{i + 1:02d}", company_id=CID, kind="employee", name=f"Staff {i + 1:02d} ({role})", country="BD",
                  source_system="synthetic" if T is T_truth else "gsheets", source_file="truth" if T is T_truth else "notebook", source_row_id=f"emp-{i + 1:02d}")
    # invoices: one per credit-sale day per household (+ opening balance invoice)
    per_hh = defaultdict(list)
    for h in hhs:
        per_hh[h.idx].append(dict(id=f"inv-{CID}-{h.idx + 1:03d}-OPEN", date=dt.date(2019, 12, 31), total=sim.meta["hh_open"][h.idx], entry=None, no="OPENING"))
    for (d, hi), amt in sorted(sim.credit_sales.items()):
        per_hh[hi].append(dict(id=f"inv-{CID}-{hi + 1:03d}-{d:%Y%m%d}", date=d, total=amt, entry=sim.credit_entry[d], no=f"B{d:%Y%m%d}-{hi + 1:03d}"))
    pays = defaultdict(list)
    for p in sim.hh_payments:
        pays[p["hh"]].append(p)
    wo = {x["hh"]: x for x in sim.writeoffs}
    pay_seq = 0
    for hi, invs in per_hh.items():
        invs.sort(key=lambda x: x["date"])
        for variant, T in (("truth", T_truth), ("books", T_books)):
            paid = {x["id"]: 0 for x in invs}
            plist = sorted(pays[hi], key=lambda p: p["date"])
            if variant == "books":
                plist = [p for p in plist if b.bid(p["entry"]) in book_ids]
            pi = 0
            rows = []
            for p in plist:
                left = p["amount"]
                for inv in invs:
                    if left <= 0:
                        break
                    if inv["date"] > p["date"] and inv["no"] != "OPENING":
                        continue
                    room = inv["total"] - paid[inv["id"]]
                    if room <= 0:
                        continue
                    take = min(room, left)
                    paid[inv["id"]] += take
                    left -= take
                    rows.append((inv, p, take))
            wrote = wo.get(hi) if (variant == "truth" or b.bid(wo[hi]["entry"] if hi in wo else "") in book_ids) else None
            for inv in invs:
                rem = inv["total"] - paid[inv["id"]]
                status = "paid" if rem == 0 else ("written_off" if wrote and inv["date"] <= wrote["date"] else ("partially_paid" if paid[inv["id"]] else "open"))
                T.add("invoices", id=inv["id"], company_id=CID, counterparty_id=f"{CID}-{hhs[hi].cid}", invoice_no=inv["no"], issue_date=inv["date"], due_date=inv["date"] + dt.timedelta(days=30),
                      currency="BDT", fx_rate="1.00000000", total=money(inv["total"]), tax="0.00", total_base=money(inv["total"]), status=status,
                      journal_entry_id=(inv["entry"] if variant == "truth" else (b.bid(inv["entry"]) if inv["entry"] else None)),
                      memo=f"Credit sale to {hhs[hi].name}", source_system="synthetic" if variant == "truth" else "gsheets",
                      source_file="truth" if variant == "truth" else f"notebook_{inv['date'].year}.xlsx", source_row_id=inv["id"])
            for inv, p, take in rows:
                pay_seq += 1
                T.add("invoice_payments", id=f"ipay-{CID}-{variant[0]}{pay_seq:06d}", company_id=CID, invoice_id=inv["id"], payment_date=p["date"], amount=money(take), currency="BDT",
                      fx_rate="1.00000000", amount_base=money(take), method=p["method"], deposit_account_id=canon.acct_id(CID, w, "wallet" if p["method"] == "wallet" else "cash"),
                      journal_entry_id=p["entry"] if variant == "truth" else b.bid(p["entry"]), source_system="synthetic" if variant == "truth" else "gsheets",
                      source_file="truth" if variant == "truth" else f"notebook_{p['date'].year}.xlsx", source_row_id=f"{p['entry']}")
    # bills
    for variant, T in (("truth", T_truth), ("books", T_books)):
        paid_by_bill = {p["bill"]: p for p in sim.bill_payments}
        for bl in sim.bills:
            p = paid_by_bill.get(bl["id"])
            T.add("bills", id=bl["id"], company_id=CID, counterparty_id=f"{CID}-sup-{bl['sup'] + 1:02d}", bill_no=bl["no"], issue_date=bl["date"], due_date=bl["due"], currency="BDT",
                  fx_rate="1.00000000", total=money(bl["total"]), tax="0.00", total_base=money(bl["total"]), status="paid" if p else "open",
                  expense_account_id=canon.acct_id(CID, w, "inv"), journal_entry_id=bl["entry"] if variant == "truth" else b.bid(bl["entry"]),
                  memo="Purchase on credit", source_system="synthetic" if variant == "truth" else "gsheets", source_file="truth" if variant == "truth" else f"notebook_{bl['date'].year}.xlsx", source_row_id=bl["id"])
            if p:
                T.add("bill_payments", id=f"bpay-{CID}-{variant[0]}{bl['no']}", company_id=CID, bill_id=bl["id"], payment_date=p["date"], amount=money(p["amount"]), currency="BDT",
                      fx_rate="1.00000000", amount_base=money(p["amount"]), method=p["method"], paid_from_account_id=canon.acct_id(CID, w, p["method"]),
                      journal_entry_id=p["entry"] if variant == "truth" else b.bid(p["entry"]), source_system="synthetic" if variant == "truth" else "gsheets",
                      source_file="truth" if variant == "truth" else f"notebook_{p['date'].year}.xlsx", source_row_id=p["entry"])
    # payroll (books == truth; the notebook shows the same salary rows)
    for variant, T in (("truth", T_truth), ("books", T_books)):
        for r in sim.payroll:
            pm = r["month"]
            T.add("payroll_runs", id=f"pr-{CID}-{variant[0]}{pm:%Y%m}-{r['idx'] + 1:02d}", company_id=CID, run_ref=f"SAL{pm:%Y%m}", period_start=pm.replace(day=1), period_end=month_end(pm),
                  pay_date=r["pay_date"], employee_id=f"{CID}-emp-{r['idx'] + 1:02d}", gross=money(r["gross"]), deductions="0.00", employer_costs="0.00", net=money(r["gross"]),
                  currency="BDT", fx_rate="1.00000000", amount_base=money(r["gross"]), paid_from_account_id=canon.acct_id(CID, w, "bank" if r["bank"] else "cash"),
                  journal_entry_id=r["entry"] if variant == "truth" else b.bid(r["entry"]), source_system="synthetic" if variant == "truth" else "gsheets",
                  source_file="truth" if variant == "truth" else f"notebook_{r['pay_date'].year}.xlsx", source_row_id=f"{pm:%Y%m}-{r['idx']}")
    # modelled BDT/USD rate (public anchor, modelled)
    for m in sorted({ym(d) for d in days()}):
        d0 = dt.date(int(m[:4]), int(m[5:]), 1)
        for T in (T_truth, T_books):
            T.add("fx_rates", id=f"fx-{CID}-USDBDT-{m}", company_id=None, rate_date=d0, from_currency="USD", to_currency="BDT", rate=f"{A.bdt_per_usd(d0):.6f}",
                  rate_source="world_bank_wdi_modelled", source_system="synthetic", source_file="data/anchors/bdt_usd_monthly_modelled.csv", source_row_id=m)
    # inventory items (categories only; no movements)
    for i, (sku, nm, unit, cat) in enumerate([("RICE", "Rice (all grades)", "kg", "merchandise"), ("OIL", "Edible oil", "litre", "merchandise"), ("SUGAR", "Sugar and flour", "kg", "merchandise"),
                                               ("LENTIL", "Lentils and pulses", "kg", "merchandise"), ("SPICE", "Spices", "kg", "merchandise"), ("DAIRY", "Dairy and eggs", "pcs", "merchandise"),
                                               ("SNACK", "Snacks and beverages", "pcs", "merchandise"), ("HOUSE", "Cleaning and household", "pcs", "merchandise")], 1):
        for T in (T_truth, T_books):
            T.add("inventory_items", id=f"{CID}-item-{sku}", company_id=CID, sku=sku, name=nm, unit=unit, category=cat, cost_method="average",
                  source_system="synthetic" if T is T_truth else "gsheets", source_file="truth" if T is T_truth else "notebook", source_row_id=sku)


def run(sim, seed, A, outdir: Path, T_truth: canon.Tables, T_books: canon.Tables):
    trace = Trace()
    rows, counts = build_statements(sim, seed)
    write_vendor_statements(sim, rows, counts, outdir, trace)
    statement_canonical(sim, rows, counts, T_truth, trace)
    write_csv(outdir / "sources" / "cash_counts.csv", ["count_date", "counted_cash_bdt"], [[c["date"].isoformat(), money(c["amount"])] for c in counts])
    write_notebook(sim, seed, outdir, trace)
    subledgers(sim, T_truth, T_books, trace, A)
    return dict(trace=trace, stmt_rows=rows, cash_counts=counts)
