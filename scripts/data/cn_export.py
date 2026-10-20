"""cn_xborder exports: Excel books, Amazon-style settlement CSVs, Wise-style statements, CNY bank CSV, canonical tables."""
from __future__ import annotations

import datetime as dt
from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook

import canon
import cn
from cn import CH, CID, WISE, FACTORIES, B2B_CUST, EMP
from core import D0, D1, Trace, add_months, days, money, month_end, rnd, rng_for, save_xlsx, write_csv, write_json, ym

BOOK_SYSTEM = "csv_upload"
BOOK_FILE = "lanxin_books_<year>.xlsx"
CH_IDX = {"US": 1, "UK": 2, "EU": 3, "STORE": 4}


def system_for(file: str) -> str:
    return "csv_upload"


def fx_at(A):
    return lambda ccy, d: 1.0 if ccy == "CNY" else A.cny_per(ccy, d)


def _d(d: dt.date) -> str:
    return d.strftime("%d-%m-%Y")


# --------------------------------------------------------------------------- statements (reality)
def build_statements(sim, A):
    w = sim.world
    rows = {"bank": [], "wise_usd": [], "wise_eur": [], "wise_gbp": []}
    conv = {c["entry"]: c for c in sim.conversions}
    ent = {e.id: e for e in w.entries}
    for e in w.entries:
        for l in e.lines:
            if l.acct not in rows:
                continue
            k = e.kind
            r = dict(date=e.date, amount=l.nat, entries=[e.id], kind=k, fee=0, ref=e.ref, cp="", desc=e.memo, ext="", exch=None)
            if l.acct.startswith("wise"):
                ccy = w.accts[l.acct].ccy
                if k == "mkt_payout":
                    ch = e.tags["channel"]
                    r.update(desc=f"Payout from {CH[ch]['mkt']} settlement", cp=("Amazon Payments" if ch != "STORE" else "Storefront payment processor"), ref=f"SETL-{ch}-{e.tags['k']:03d}")
                elif k == "b2b_payment":
                    r.update(desc="Payment received", cp=e.memo.split(" for ")[0].replace("Payment from ", ""), ref=e.ref.replace("WSPAY-", ""))
                elif k == "conversion":
                    cv = conv[e.id]
                    r.update(desc=f"Converted {ccy} to CNY", cp="Wise Asia-Pacific (CNY payout to bank)", fee=rnd(cv["nat"] * 0.0055), exch=(ccy, "CNY", A.cny_per(ccy, e.date), cv["cny"]))
                elif k == "freight":
                    r.update(desc="Transfer to freight forwarder", cp="Pacific Star Logistics (HK) Ltd", ref="freight invoice")
                elif k == "duty":
                    r.update(desc="Transfer to customs broker", cp="Shenzhen Customs Broker Co.", ref="duties and clearance")
                elif k == "store_ads":
                    r.update(desc="Card payment", cp="Meta Platforms Ireland Ltd", ref="ads")
                elif k == "opening":
                    r.update(desc="Opening balance brought forward", cp="", ref="OPENING")
            else:
                if k == "conversion":
                    r.update(desc="WISE ASIA-PACIFIC PAYOUT", cp="Wise", ref=e.ref)
                elif k in ("po_deposit", "po_receipt"):
                    fac = next(x for x in sim.purchases if x["po"] == e.ref.rsplit("-", 1)[0])["fac"]
                    r.update(desc=("DEPOSIT 30% " if k == "po_deposit" else "BALANCE 70% ") + e.ref, cp=FACTORIES[fac])
                elif k in ("payroll", "payroll_social"):
                    r.update(desc="SALARY PAYMENT BATCH" if k == "payroll" else "SOCIAL INSURANCE AND HOUSING FUND", cp="Bank payroll batch" if k == "payroll" else "Tax bureau / housing fund centre")
                elif k == "rent":
                    r.update(desc="RENT", cp="Longhua Industrial Park Property Co.")
                elif k == "software":
                    r.update(desc="SERVICE FEE", cp="Software subscriptions")
                elif k == "accounting":
                    r.update(desc="BOOKKEEPING FEE", cp="Shenzhen Qiancheng Accounting Service")
                elif k == "bank_fee":
                    r.update(desc="ACCOUNT AND REMITTANCE CHARGES", cp="Bank")
                elif k == "cash_withdrawal":
                    r.update(desc="CASH WITHDRAWAL", cp="Bank counter")
                elif k in ("loan_drawdown", "loan_repayment", "loan_interest"):
                    r.update(desc={"loan_drawdown": "LOAN DISBURSEMENT", "loan_repayment": "LOAN REPAYMENT", "loan_interest": "LOAN INTEREST"}[k], cp="Working-capital lender")
                elif k == "drawings":
                    r.update(desc="DIVIDEND PAYMENT", cp="Owner")
                elif k == "owner_injection":
                    r.update(desc="CAPITAL TRANSFER IN", cp="Owner")
                elif k == "opening":
                    r.update(desc="Opening balance brought forward", ref="OPENING")
            rows[l.acct].append(r)
    for a, rs in rows.items():
        rs.sort(key=lambda r: (r["date"], r["entries"][0]))
        bal = 0
        for i, r in enumerate(rs):
            bal += r["amount"]
            r["balance"] = bal
            r["ext"] = f"{a[:2].upper()}{r['date']:%y%m%d}{i % 100000:05d}"
        assert bal == w.bal(a), (a, bal, w.bal(a))
    counts = []
    bal = 0
    day_amt = defaultdict(int)
    for e in w.entries:
        for l in e.lines:
            if l.acct == "petty":
                day_amt[e.date] += l.nat
    for d in days():
        bal += day_amt.get(d, 0)
        if d == month_end(d):
            counts.append(dict(date=d, amount=bal))
    return rows, counts


WISE_HDR = ["TransferWise ID", "Date", "Amount", "Currency", "Description", "Payment Reference", "Running Balance", "Exchange From", "Exchange To", "Exchange Rate",
            "Payer Name", "Payee Name", "Payee Account Number", "Merchant", "Card Last Four Digits", "Card Holder Full Name", "Attachment", "Note", "Total fees", "Exchange To Amount"]


def write_statements(sim, rows, counts, outdir: Path, trace: Trace):
    src = outdir / "sources"
    w = sim.world
    for acct in ("wise_usd", "wise_eur", "wise_gbp"):
        ccy = w.accts[acct].ccy
        fn = f"wise_style_statement_{ccy}.csv"
        out = []
        for i, r in enumerate(rows[acct], 2):
            a = r["amount"]
            payer = r["cp"] if a > 0 else ""
            payee = r["cp"] if a < 0 else ""
            merchant = r["cp"] if r["kind"] == "store_ads" else ""
            ex = r["exch"]
            out.append([r["ext"], _d(r["date"]), money(a), ccy, r["desc"], r["ref"], money(r["balance"]), ex[0] if ex else "", ex[1] if ex else "", f"{ex[2]:.6f}" if ex else "",
                        payer, payee, "", merchant, "4417" if merchant else "", "L. Zhang" if merchant else "", "", "", money(r["fee"]) if r["fee"] else "0.00", money(ex[3]) if ex else ""])
            r["src"] = (fn, str(i))
            for eid in r["entries"]:
                trace.add(fn, str(i), "truth", eid)
        write_csv(src / fn, WISE_HDR, out)
    by_year = defaultdict(list)
    for r in rows["bank"]:
        by_year[r["date"].year].append(r)
    for y, rs in sorted(by_year.items()):
        fn = f"cny_bank_statement_{y}.csv"
        out = []
        for i, r in enumerate(rs, 2):
            a = r["amount"]
            out.append([r["date"].isoformat(), r["date"].isoformat(), r["cp"], r["desc"], r["ref"], money(-a) if a < 0 else "", money(a) if a > 0 else "", money(r["balance"])])
            r["src"] = (fn, str(i))
            for eid in r["entries"]:
                trace.add(fn, str(i), "truth", eid)
        write_csv(src / fn, ["Transaction Date", "Value Date", "Counterparty", "Summary", "Reference", "Debit (CNY)", "Credit (CNY)", "Balance (CNY)"], out)
    write_csv(src / "petty_cash_counts.csv", ["count_date", "counted_cash_cny"], [[c["date"].isoformat(), money(c["amount"])] for c in counts])
    for ci, c in enumerate(counts, 2):
        trace.add("petty_cash_counts.csv", str(ci), "truth_balance", f"petty:{c['date']}")


def statement_canonical(sim, rows, counts, T: canon.Tables):
    w = sim.world
    n = 0
    for acct in ("bank", "wise_usd", "wise_eur", "wise_gbp"):
        a = w.accts[acct]
        for r in rows[acct]:
            n += 1
            f, rid = r["src"]
            fxr = 1.0 if a.ccy == "CNY" else sim.world.fx(a.ccy, r["date"])
            T.add("external_statement_lines", id=f"ext-{CID}-{n:06d}", company_id=CID, channel="bank", account_id=canon.acct_id(CID, w, acct),
                  statement_ref="****6608" if acct == "bank" else f"****2217 {a.ccy}", txn_date=r["date"], value_date=r["date"], amount=money(r["amount"]), currency=a.ccy,
                  fx_rate=f"{fxr:.8f}", amount_base=money(rnd(r["amount"] * fxr)), description=r["desc"], counterparty_text=r["cp"] or None, balance=money(r["balance"]),
                  external_txn_id=r["ext"], fee=money(r["fee"]), source_system="bank_csv" if acct == "bank" else "wise_csv", source_file=f, source_row_id=rid)
    for ci, c in enumerate(counts, 2):
        n += 1
        T.add("external_statement_lines", id=f"ext-{CID}-{n:06d}", company_id=CID, channel="cash_count", account_id=canon.acct_id(CID, w, "petty"), statement_ref="office safe",
              txn_date=c["date"], value_date=c["date"], amount=money(c["amount"]), currency="CNY", fx_rate="1.00000000", amount_base=money(c["amount"]),
              description="Month-end petty cash count", balance=money(c["amount"]), external_txn_id=f"COUNT{c['date']:%Y%m%d}", fee="0.00",
              source_system="cash_count", source_file="petty_cash_counts.csv", source_row_id=str(ci))


# --------------------------------------------------------------------------- marketplace settlement files
SET_HDR = ["settlement-id", "settlement-start-date", "settlement-end-date", "deposit-date", "total-amount", "currency", "transaction-type", "order-id", "merchant-order-id",
           "adjustment-id", "shipment-id", "marketplace-name", "amount-type", "amount-description", "amount", "fulfillment-id", "posted-date", "posted-date-time",
           "order-item-code", "merchant-order-item-id", "merchant-adjustment-item-id", "sku", "quantity-purchased", "promotion-id"]


def settlement_id(c: str, k: int) -> str:
    return str(20_000_000_000 + CH_IDX[c] * 1000 + k)


def write_settlements(sim, outdir: Path, trace: Trace, T=None):
    M = sim.meta
    D = sim.daily
    per_file: dict[tuple, list] = defaultdict(list)
    file_ch: dict[str, str] = {}
    index = []
    for c in CH:
        p = CH[c]
        for k, (a, b) in sorted(M["periods"].items()):
            last = min(b, D1)
            sid = settlement_id(c, k)
            pid = M["pay_ids"].get((c, k))
            info = M["pay_info"].get(pid) if pid else None
            closed = b + dt.timedelta(days=2) <= D1 and pid is not None
            dep = (b + dt.timedelta(days=2)) if closed else None
            total = sim.world_ent[pid].lines[0].dr if closed else None
            fn = f"amazon_style_settlement_{c}_{a.year}.csv" if c != "STORE" else f"storefront_payout_report_{a.year}.csv"
            rows = per_file[(fn)]
            file_ch[fn] = c
            def ts(d, end=False):
                return f"{d.isoformat()} {'23:59:59' if end else '00:00:00'} UTC"
            refs_sale = M["sale_ids"].get((c, k), [])
            refs_fee = M["fee_ids"].get((c, k), [])
            refs_ref = M["refund_ids"].get((c, k), [])
            refs_reim = M["reimb_ids"].get((c, k), [])
            summary_refs = ([pid] if pid else []) + refs_sale
            start_row = len(rows)
            rows.append(([sid, ts(a), ts(b, True), ts(dep, True) if dep else "", money(total) if total is not None else "", p["ccy"]] + [""] * 18, summary_refs))
            net = 0
            def line(d, ttype, atype, adesc, amt, refs, oid=None):
                nonlocal net
                if not amt:
                    return
                net += amt
                rows.append(([sid, "", "", "", "", p["ccy"], ttype, oid or f"AGG-{d:%Y%m%d}", "", "", "", p["mkt"], atype, adesc, money(amt), "", d.isoformat(), f"{d.isoformat()}T12:00:00+00:00",
                              "", "", "", "", "", ""], refs))
            for d in days(a, last):
                r = D[(c, d)]
                line(d, "Order", "ItemPrice", "Principal", r["prin"], refs_sale)
                if c != "STORE":
                    line(d, "Order", "ItemFees", "Commission", -(r["ref"] - r["ref_refund_fee"]), refs_fee)
                    line(d, "Order", "ItemFees", "FBAPerUnitFulfillmentFee", -r["fba"], refs_fee)
                    line(d, "Refund", "ItemPrice", "Principal", -r["refunds"], refs_ref)
                    line(d, "Refund", "ItemFees", "RefundCommission", -r["ref_refund_fee"], refs_fee)
                    line(d, "other-transaction", "other-transaction", "Cost of Advertising", -r["ads"], refs_fee)
                    line(d, "other-transaction", "other-transaction", "Storage Fee", -r.get("storage", 0), refs_fee)
                    line(d, "other-transaction", "other-transaction", "Subscription Fee", -r.get("sub", 0), refs_fee)
                    line(d, "other-transaction", "other-transaction", "REVERSAL_REIMBURSEMENT", r["reimb"], refs_reim)
                else:
                    line(d, "Order", "ItemFees", "Payment processing fee", -r["pay"], refs_fee)
                    line(d, "Refund", "ItemPrice", "Principal", -r["refunds"], refs_ref)
            if info:
                if info["reserve_old"]:
                    line(last, "other-transaction", "other-transaction", "Previous Reserve Amount Balance", info["reserve_old"], [pid], oid=f"RESERVE-{k:03d}")
                if info["reserve_new"]:
                    line(last, "other-transaction", "other-transaction", "Current Reserve Amount", -info["reserve_new"], [pid], oid=f"RESERVE-{k:03d}")
            if closed:
                assert net == total, (c, k, net, total)
            index.append(dict(settlement_id=sid, channel=c, marketplace=p["mkt"], currency=p["ccy"], start=a.isoformat(), end=b.isoformat(), deposit_date=dep.isoformat() if dep else None,
                              total_deposit=money(total) if total is not None else None, status="closed" if closed else "open", file=fn, payout_entry=pid))
    for fn, rows in per_file.items():
        out = [r[0] for r in rows]
        write_csv(outdir / "sources" / fn, SET_HDR, out)
        c = file_ch[fn]
        ccy = CH[c]["ccy"]
        for i, (r, refs) in enumerate(rows, 2):
            for ref in refs:
                trace.add(fn, str(i), "truth", ref)
            if T is None:
                continue
            if r[6] == "":  # settlement summary row -> payout line (skipped while the settlement is open)
                if r[4] == "":
                    continue
                dd = dt.date.fromisoformat(r[3][:10])
                amt, desc = -round(float(r[4]) * 100), "Payout to bank account"
            else:
                dd = dt.date.fromisoformat(r[16])
                amt, desc = round(float(r[14]) * 100), f"{r[6]}: {r[13]}"
            nset = len(T.t["external_statement_lines"]) + 1
            T.add("external_statement_lines", id=f"ext-{CID}-S{nset:06d}", company_id=CID, channel="pos_settlement", account_id=canon.acct_id(CID, sim.world, CH[c]["clr"]),
                  statement_ref=r[0], txn_date=dd, value_date=dd, amount=money(amt), currency=ccy, fx_rate=f"{sim.world.fx(ccy, dd):.8f}", amount_base=money(rnd(amt * sim.world.fx(ccy, dd))),
                  description=desc, external_txn_id=r[0], fee="0.00", source_system="marketplace_settlement", source_file=fn, source_row_id=str(i))
    return index


# --------------------------------------------------------------------------- Excel books
def write_books_xlsx(sim, A, outdir: Path, trace: Trace):
    b = sim.book
    w = sim.world
    book_ids = set(b.entries)
    by_year = defaultdict(list)
    for e in b.sorted_entries():
        by_year[e.date.year].append(e)
    ent = {e.id: e for e in w.entries}
    for y in range(2020, 2027):
        wb = Workbook()
        ws = wb.active
        ws.title = "GL_Journal"
        fn = f"lanxin_books_{y}.xlsx"
        ws.append(["Voucher", "Date", "Account code", "Account name", "Currency", "Debit", "Credit", "Rate to CNY", "Memo"])
        rowno = 1
        for e in by_year[y]:
            seen_row = False
            for l in e.lines:
                a = w.accts[l.acct]
                rowno += 1
                ws.append([e.id.split("-")[-1] if "-A" not in e.id else "A" + e.id.split("-A")[-1], dt.datetime(e.date.year, e.date.month, e.date.day), a.code, a.name, l.ccy,
                           (l.dr / 100) or None, (l.cr / 100) or None, l.fx if l.ccy != "CNY" else 1, e.memo])
                trace.add(fn, f"GL_Journal!R{rowno}", "book", e.id)
        ws2 = wb.create_sheet("Sales_by_Channel")
        ws2.append(["Month", "Channel", "Currency", "Sales (local)", "Sales (CNY at booked rate)"])
        r2 = 1
        agg = defaultdict(lambda: [0, 0, []])
        for e in by_year[y]:
            for l in e.lines:
                if l.acct.startswith("rev_") and l.cr:
                    x = agg[(ym(e.date), l.acct)]
                    x[0] += l.cr
                    x[1] += l.base
                    x[2].append(e.id)
        for (m, acct), (loc, base_, ids) in sorted(agg.items()):
            r2 += 1
            ws2.append([m, w.accts[acct].name, "EUR" if acct == "rev_b2b" else CH.get({"rev_us": "US", "rev_uk": "UK", "rev_eu": "EU", "rev_store": "STORE"}[acct], {"ccy": "USD"})["ccy"], loc / 100, -base_ / 100])
            for i in ids:
                trace.add(fn, f"Sales_by_Channel!R{r2}", "book", i)
        ws3 = wb.create_sheet("PO_Tracker")
        ws3.append(["PO", "Supplier", "Order date", "Goods received", "Total (CNY)", "Deposit 30% (CNY)", "Balance 70% (CNY)", "Status"])
        r3 = 1
        for po in sim.purchases:
            if po["order_date"].year == y:
                r3 += 1
                rec = po.get("receipt_entry")
                ws3.append([po["po"], FACTORIES[po["fac"]], dt.datetime(*po["order_date"].timetuple()[:3]), dt.datetime(*po["land"].timetuple()[:3]) if rec else None, po["total"] / 100, po["dep"] / 100,
                            (po["total"] - po["dep"]) / 100, "received" if rec else "in production"])
                trace.add(fn, f"PO_Tracker!R{r3}", "truth", po["dep_entry"])
                if rec:
                    trace.add(fn, f"PO_Tracker!R{r3}", "truth", rec)
        ws4 = wb.create_sheet("Supplier_Payments")
        ws4.append(["Date", "Reference", "Supplier", "Type", "Amount (CNY)"])
        r4 = 1
        for e in by_year[y]:
            if e.kind in ("po_deposit", "po_receipt", "dup_deposit"):
                r4 += 1
                po = e.ref.rsplit("-", 1)[0]
                p_ = next((x for x in sim.purchases if x["po"] == po), None)
                amt = next(l.cr for l in e.lines if l.acct == "bank") if e.kind != "dup_deposit" else e.lines[1].cr
                ws4.append([dt.datetime(e.date.year, e.date.month, e.date.day), e.ref, FACTORIES[p_["fac"]] if p_ else "", "Deposit" if "DEP" in e.ref else "Balance", amt / 100])
                trace.add(fn, f"Supplier_Payments!R{r4}", "book", e.id)
        ws5 = wb.create_sheet("AR_Wholesale")
        ws5.append(["Invoice", "Customer", "Invoice date", "Due date", "Amount (EUR)", "Paid?"])
        r5 = 1
        for r in sim.b2b:
            if r["date"].year == y:
                r5 += 1
                paid = r.get("paid") and b.bid(r["pay_entry"]) in book_ids
                ws5.append([r["no"], B2B_CUST[r["cust"]], dt.datetime(*r["date"].timetuple()[:3]), dt.datetime(*r["due"].timetuple()[:3]), r["total"] / 100, "Yes" if paid else "No"])
                trace.add(fn, f"AR_Wholesale!R{r5}", "truth", r["entry"])
        save_xlsx(wb, outdir / "sources" / fn)


# --------------------------------------------------------------------------- canonical subledgers
def subledgers(sim, T_truth, T_books, A):
    w, b = sim.world, sim.book
    book_ids = set(b.entries)
    for T, ss, sf in ((T_truth, "synthetic", "truth"), (T_books, "csv_upload", "lanxin_books")):
        for i, f in enumerate(FACTORIES, 1):
            T.add("counterparties", id=f"{CID}-sup-{i:02d}", company_id=CID, kind="supplier", name=f, country="CN", payment_terms_days=30, external_ref=f"sup-{i:02d}", source_system=ss, source_file=sf, source_row_id=f"sup-{i:02d}")
        for i, f in enumerate(B2B_CUST, 1):
            T.add("counterparties", id=f"{CID}-cust-{i:02d}", company_id=CID, kind="customer", name=f, country="DE", payment_terms_days=45, external_ref=f"cust-{i:02d}", source_system=ss, source_file=sf, source_row_id=f"cust-{i:02d}")
        for i, (role, _) in enumerate(EMP, 1):
            T.add("counterparties", id=f"{CID}-emp-{i:02d}", company_id=CID, kind="employee", name=f"Staff {i:02d} ({role})", country="CN", source_system=ss, source_file=sf, source_row_id=f"emp-{i:02d}")
        for it in ("cables", "chargers", "cases", "stands", "packaging"):
            T.add("inventory_items", id=f"{CID}-item-{it}", company_id=CID, sku=it.upper(), name=it.capitalize(), unit="pcs", category="finished_good", cost_method="average", source_system=ss, source_file=sf, source_row_id=it)
    # wholesale invoices
    for variant, T in (("truth", T_truth), ("books", T_books)):
        for r in sim.b2b:
            paid_in_book = r.get("paid") and (variant == "truth" or b.bid(r["pay_entry"]) in book_ids)
            wo = r.get("written_off") and (variant == "truth" or b.bid(r["wo_entry"]) in book_ids)
            status = "paid" if paid_in_book else ("written_off" if wo else "open")
            fx = w.fx("EUR", r["date"])
            T.add("invoices", id=r["id"], company_id=CID, counterparty_id=f"{CID}-cust-{r['cust'] + 1:02d}", invoice_no=r["no"], issue_date=r["date"], due_date=r["due"], currency="EUR", fx_rate=f"{fx:.8f}",
                  total=money(r["total"]), tax="0.00", total_base=money(rnd(r["total"] * fx)), status=status, journal_entry_id=r["entry"] if variant == "truth" else b.bid(r["entry"]),
                  memo="EU wholesale invoice", source_system="synthetic" if variant == "truth" else "csv_upload", source_file="truth" if variant == "truth" else f"lanxin_books_{r['date'].year}.xlsx", source_row_id=r["no"])
            if paid_in_book:
                fx2 = w.fx("EUR", r["pay_actual"])
                T.add("invoice_payments", id=f"ipay-{CID}-{variant[0]}{r['no']}", company_id=CID, invoice_id=r["id"], payment_date=r["pay_actual"], amount=money(r["total"]), currency="EUR", fx_rate=f"{fx2:.8f}",
                      amount_base=money(rnd(r["total"] * fx2)), method="bank", deposit_account_id=canon.acct_id(CID, w, "wise_eur"), journal_entry_id=r["pay_entry"] if variant == "truth" else b.bid(r["pay_entry"]),
                      source_system="synthetic" if variant == "truth" else "csv_upload", source_file="truth" if variant == "truth" else f"lanxin_books_{r['pay_actual'].year}.xlsx", source_row_id=r["pay_entry"])
    # factory bills + payments
    for variant, T in (("truth", T_truth), ("books", T_books)):
        for po in sim.purchases:
            if "receipt_entry" not in po:
                continue
            bid_ = f"bill-{CID}-{po['po']}"
            T.add("bills", id=bid_, company_id=CID, counterparty_id=f"{CID}-sup-{po['fac'] + 1:02d}", bill_no=po["po"], issue_date=po["order_date"], due_date=po["land"], currency="CNY", fx_rate="1.00000000",
                  total=money(po["total"]), tax="0.00", total_base=money(po["total"]), status="paid", expense_account_id=canon.acct_id(CID, w, "inv"),
                  journal_entry_id=po["receipt_entry"] if variant == "truth" else b.bid(po["receipt_entry"]), memo="Factory purchase order", source_system="synthetic" if variant == "truth" else "csv_upload",
                  source_file="truth" if variant == "truth" else f"lanxin_books_{po['order_date'].year}.xlsx", source_row_id=po["po"])
            for kind_, amt, dd, en in (("deposit", po["dep"], po["order_date"], po["dep_entry"]), ("balance", po["total"] - po["dep"], po["land"], po["receipt_entry"])):
                T.add("bill_payments", id=f"bpay-{CID}-{variant[0]}{po['po']}-{kind_[0]}", company_id=CID, bill_id=bid_, payment_date=dd, amount=money(amt), currency="CNY", fx_rate="1.00000000",
                      amount_base=money(amt), method="bank", paid_from_account_id=canon.acct_id(CID, w, "bank"), journal_entry_id=en if variant == "truth" else b.bid(en),
                      source_system="synthetic" if variant == "truth" else "csv_upload", source_file="truth" if variant == "truth" else f"lanxin_books_{dd.year}.xlsx", source_row_id=f"{po['po']}-{kind_}")
    # payroll
    for variant, T in (("truth", T_truth), ("books", T_books)):
        for pr in sim.payroll:
            n = len(pr["rows"])
            soc = weighted_even(pr["social"], n)
            for (i, role, gross), s in zip(pr["rows"], soc):
                pm = pr["month"]
                T.add("payroll_runs", id=f"pr-{CID}-{variant[0]}{pm:%Y%m}-{i + 1:02d}", company_id=CID, run_ref=f"PAY{pm:%Y%m}", period_start=pm.replace(day=1), period_end=month_end(pm), pay_date=pr["pay_date"],
                      employee_id=f"{CID}-emp-{i + 1:02d}", gross=money(gross), deductions="0.00", employer_costs=money(s), net=money(gross), currency="CNY", fx_rate="1.00000000",
                      amount_base=money(gross + s), paid_from_account_id=canon.acct_id(CID, w, "bank"), journal_entry_id=pr["entry"] if variant == "truth" else b.bid(pr["entry"]),
                      source_system="synthetic" if variant == "truth" else "csv_upload", source_file="truth" if variant == "truth" else f"lanxin_books_{pr['pay_date'].year}.xlsx", source_row_id=f"{pm:%Y%m}-{i}")
    # public FX series (global rows, company_id NULL): weekly to keep size small
    d = D0
    while d <= D1:
        for ccy in ("USD", "EUR", "GBP"):
            for T in (T_truth, T_books):
                T.add("fx_rates", id=f"fx-{ccy}CNY-{d:%Y%m%d}", company_id=None, rate_date=d, from_currency=ccy, to_currency="CNY", rate=f"{A.cny_per(ccy, d):.6f}",
                      rate_source="fed_h10_via_fred", source_system="synthetic", source_file="data/anchors/fx_h10_daily.csv", source_row_id=f"{d.isoformat()}-{ccy}")
        d += dt.timedelta(days=7)


def weighted_even(total: int, n: int) -> list[int]:
    base, rem = divmod(total, n)
    return [base + (1 if i < rem else 0) for i in range(n)]


def stale_findings(sim, asofs):
    """Open wholesale invoices more than 90 days past due at each as-of date (a finding, not a ledger difference)."""
    out = []
    n = 0
    for asof in asofs:
        for r in sim.b2b:
            settled = r.get("pay_actual") if r.get("paid") else None
            if r["date"] > asof:
                continue
            if settled and settled <= asof:
                continue
            if r.get("written_off") and r["written_off"] <= asof:
                continue
            days_late = (asof - r["due"]).days
            if days_late > 90:
                n += 1
                out.append(dict(id=f"{CID}-S{n:05d}", type="stale_receivable", subtype="overdue_over_90_days", date=asof.isoformat(), as_of=asof.isoformat(), amount=money(r["total"]), currency="EUR",
                                amount_base=money(rnd(r["total"] * sim.world.fx("EUR", asof))), description=f"{B2B_CUST[r['cust']]} invoice {r['no']} is {days_late} days past due at {asof}",
                                truth_entries=[r["entry"]], book_entries=[], affected_source_rows=[], fix={"void_book_entries": [], "add_entries": [], "note": "no ledger change; chase or provide for the debt"},
                                details=dict(invoice=r["id"], days_past_due=days_late, eventually_paid=bool(r.get("paid")), never_collected=bool(r.get("never")))))
    return out


def run(sim, seed, A, outdir: Path, T_truth, T_books):
    trace = Trace()
    sim.world_ent = {e.id: e for e in sim.world.entries}
    rows, counts = build_statements(sim, A)
    write_statements(sim, rows, counts, outdir, trace)
    statement_canonical(sim, rows, counts, T_truth)
    index = write_settlements(sim, outdir, trace, T_truth)
    write_books_xlsx(sim, A, outdir, trace)
    subledgers(sim, T_truth, T_books, A)
    from key import asof_dates
    findings = stale_findings(sim, asof_dates())
    write_json(outdir / "truth" / "settlement_index.json", index)
    return dict(trace=trace, stmt_rows=rows, key_extra=dict(marketplace_settlements=index), findings=findings)
