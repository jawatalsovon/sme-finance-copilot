"""Self-check of data/synthetic/ : python scripts/data/validate.py  (exit code 1 on any failure)."""
from __future__ import annotations

import csv
import gzip
import json
import sys
from collections import Counter, defaultdict
from decimal import Decimal
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
SYN = ROOT / "data" / "synthetic"
COMPANIES = ["bd_grocer", "cn_xborder", "uk_garden"]
csv.field_size_limit(10 ** 8)
FAILS: list[str] = []


def check(ok: bool, msg: str) -> None:
    print(("  PASS " if ok else "  FAIL ") + msg)
    if not ok:
        FAILS.append(msg)


def rd(path: Path):
    for p in (path, Path(str(path) + ".gz")):
        if p.exists():
            f = gzip.open(p, "rt", encoding="utf-8", newline="") if p.suffix == ".gz" else open(p, encoding="utf-8", newline="")
            with f:
                yield from csv.DictReader(f)
            return
    raise FileNotFoundError(path)


def D(x: str) -> Decimal:
    return Decimal(x) if x not in ("", None) else Decimal(0)


def load_journal(folder: Path):
    lines = defaultdict(list)
    dates = {}
    for e in rd(folder / "journal_entries.csv"):
        dates[e["id"]] = e["entry_date"]
    for l in rd(folder / "journal_lines.csv"):
        lines[l["entry_id"]].append(l)
    return dates, lines


def sig(date, l):
    return (date, l["account_id"], l["currency"], str(D(l["debit"]) - D(l["credit"])), str(D(l["amount_base"])))


def validate(cid: str) -> None:
    print(f"== {cid}")
    root = SYN / cid
    key = json.loads((root / "answer_key.json").read_text(encoding="utf-8"))
    accts = {a["id"]: a for a in rd(root / "truth" / "accounts.csv")}
    t_dates, t_lines = load_journal(root / "truth")
    b_dates, b_lines = load_journal(root / "books")
    # 1 balance
    bad_t = [e for e, ls in t_lines.items() if sum(D(l["amount_base"]) for l in ls) != 0]
    bad_b = [e for e, ls in b_lines.items() if sum(D(l["amount_base"]) for l in ls) != 0]
    check(not bad_t, f"truth: all {len(t_lines)} journals balance (debits = credits in base currency)")
    check(not bad_b, f"books: all {len(b_lines)} journals balance")
    bad_sign = [l for ls in t_lines.values() for l in ls if D(l["debit"]) < 0 or D(l["credit"]) < 0]
    check(not bad_sign, "truth: debit/credit never negative")
    # 2 P&L reproduces
    pnl = defaultdict(lambda: defaultdict(Decimal))
    for e, ls in t_lines.items():
        m = t_dates[e][:7]
        for l in ls:
            a = accts[l["account_id"]]
            v = D(l["amount_base"])
            if a["type"] == "income":
                pnl[m]["income"] += -v
                if a["subtype"] == "revenue":
                    pnl[m]["revenue"] += -v
            elif a["type"] == "expense":
                pnl[m]["expense"] += v
    mism = []
    for r in key["true_monthly_pnl"]:
        p = pnl[r["month"]]
        if D(r["revenue"]) != p["revenue"] or D(r["net_profit"]) != p["income"] - p["expense"]:
            mism.append(r["month"])
    check(not mism, f"truth monthly P&L in answer key reproduces from truth journals to the unit ({len(key['true_monthly_pnl'])} months)")
    # 3 coverage
    months = {f"{y}-{m:02d}" for y in range(2020, 2027) for m in range(1, 13)} - {f"2026-{m:02d}" for m in (10, 11, 12)}
    check(months <= {d[:7] for d in t_dates.values()}, "truth covers every month 2020-01..2026-09 with journal entries")
    check(months <= {d[:7] for d in b_dates.values()}, "books cover every month 2020-01..2026-09")
    check(all(float(pnl[m]["revenue"]) > 0 for m in months), "no month with zero revenue (truth)")
    # 4 books + fixes = truth
    book_sig = Counter()
    for e, ls in b_lines.items():
        for l in ls:
            book_sig[sig(b_dates[e], l)] += 1
    code2id = {a["code"]: a["id"] for a in accts.values()}
    cur = {a["id"]: a["currency"] for a in accts.values()}
    # apply fixes
    for p in key["problems"]:
        for bid in p["fix"]["void_book_entries"]:
            for l in b_lines[bid]:
                book_sig[sig(b_dates[bid], l)] -= 1
        for ent in p["fix"]["add_entries"]:
            for l in ent["lines"]:
                book_sig[(ent["date"], l["account"], l["currency"], str(D(l["debit"]) - D(l["credit"])), str(D(l["amount_base"])))] += 1
    truth_sig = Counter()
    for e, ls in t_lines.items():
        for l in ls:
            truth_sig[sig(t_dates[e], l)] += 1
    conv = Counter({k: n for k, n in book_sig.items() if n})
    tconv2 = Counter(truth_sig)
    diff = (conv - tconv2) + (tconv2 - conv)
    check(not diff, f"books + answer-key fixes == truth exactly (same date/account/amount lines); differing lines: {len(diff)}")
    # 5 gap closes
    for d, g in key["cash_truth_gap"].items():
        if D(g["unexplained"]) != 0:
            check(False, f"gap at {d} fully explained by planted problems")
            break
    else:
        check(True, f"Cash Truth gap fully explained by planted problems at all {len(key['cash_truth_gap'])} as-of dates")
    # 6 statements = truth
    last = {}
    for r in rd(root / "truth" / "external_statement_lines.csv"):
        if r["channel"] in ("bank", "wallet"):
            last[r["account_id"]] = (r["txn_date"], D(r["balance"]), r["currency"])
    me = key["true_month_end_cash"][-1]["accounts"]
    ok = True
    for a in me:
        code = a["account"].split()[0]
        aid = f"{cid}-{code}"
        if aid in last and last[aid][1] != D(a["true_balance"]):
            ok = False
    check(ok, "last statement balance per bank/wallet account equals the true 2026-09-30 balance")
    # 7 trace
    trace = defaultdict(set)
    refs = set()
    for r in rd(root / "truth" / "source_trace.csv"):
        trace[r["file"]].add(r["row_id"])
        refs.add((r["ref_kind"], r["ref_id"]))
    src = root / "sources"
    miss = []
    nfiles = 0
    for f in sorted(src.iterdir()):
        if f.suffix == ".json":
            continue
        nfiles += 1
        if f.suffix == ".csv":
            with open(f, encoding="utf-8", newline="") as h:
                n = sum(1 for _ in h) - 1
            have = {str(i) for i in range(2, n + 2)}
            miss += [(f.name, x) for x in sorted(have - trace[f.name])[:3]]
        else:
            from openpyxl import load_workbook
            wb = load_workbook(f, read_only=True)
            for ws in wb.worksheets:
                hdr = None
                for i, row in enumerate(ws.iter_rows(values_only=True), 1):
                    first = str(row[0]) if row and row[0] is not None else ""
                    if hdr is None:
                        if first in ("Date", "Voucher", "Month", "PO", "Invoice"):
                            hdr = i
                        continue
                    if any(c is not None for c in row) and f"{ws.title}!R{i}" not in trace[f.name]:
                        miss.append((f.name, f"{ws.title}!R{i}"))
    check(not miss, f"every source-export row ({nfiles} files) traces to a truth/book entry id; untraced sample: {miss[:3]}")
    ids = set(t_lines) | set(b_lines)
    badref = [r for r in refs if r[0] in ("truth", "book") and r[1] not in ids]
    check(not badref, "trace references point at existing entries")
    # 8 problems reference real entries
    bad = [p["id"] for p in key["problems"] if any(t not in t_lines for t in p["truth_entries"]) or any(b not in b_lines for b in p["book_entries"])]
    check(not bad, f"all {len(key['problems'])} problem records reference existing entries")
    print("  planted:", key["summary"]["planted_problem_counts"])


def id2k(aid, km):
    return km.get(aid, aid)


def sizes() -> None:
    print("== sizes")
    total = 0
    for p in SYN.rglob("*"):
        if p.is_file():
            s = p.stat().st_size
            total += s
            if s > 50 * 1024 * 1024:
                print(f"  WARN single file > 50 MB: {p}")
    check(total < 150 * 1024 * 1024, f"total output {total / 1e6:.1f} MB (< 150 MB)")


if __name__ == "__main__":
    for c in (sys.argv[1:] or COMPANIES):
        validate(c)
    sizes()
    print("\nALL CHECKS PASSED" if not FAILS else f"\n{len(FAILS)} CHECK(S) FAILED")
    sys.exit(1 if FAILS else 0)
